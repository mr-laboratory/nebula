"""Word (.docx) export: real Word styles, so headings and lists stay editable."""

from typing import IO

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from markdown_it.token import Token

from app.rendering.document import ExportDocument, ExportPost, parse, safe_href, status_line

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"
LINK_COLOR = RGBColor(0x43, 0x38, 0xCA)
MUTED_COLOR = RGBColor(0x6B, 0x72, 0x80)
MAX_HEADING = 9


def render(document: ExportDocument, target: IO[bytes]) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = BODY_FONT
    doc.styles["Normal"].font.size = Pt(11)
    _set_properties(doc, document)
    if document.collection:
        _cover(doc, document)
    for index, post in enumerate(document.posts):
        if index or document.collection:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _post(doc, post, anchor=f"post-{index + 1}")
    doc.save(target)


def _set_properties(doc: DocxDocument, document: ExportDocument) -> None:
    # Replace the template's defaults ("python-docx") so the file describes itself honestly.
    props = doc.core_properties
    props.title = document.title
    props.author = document.author
    props.last_modified_by = document.author
    props.comments = ""
    props.subject = ""
    props.keywords = ""
    props.created = props.modified = document.generated_at


def _cover(doc: DocxDocument, document: ExportDocument) -> None:
    doc.add_paragraph(document.title, style="Title")
    meta = doc.add_paragraph()
    _muted(meta.add_run(f"{document.author} · exported {document.generated_at:%d %B %Y}"))
    doc.add_heading("Contents", level=2)
    for index, post in enumerate(document.posts):
        entry = doc.add_paragraph(style="List Number")
        _internal_link(entry, f"post-{index + 1}", post.title)


def _post(doc: DocxDocument, post: ExportPost, anchor: str) -> None:
    heading = doc.add_heading(level=1)
    _bookmark(heading, anchor, post.title)
    _muted(doc.add_paragraph().add_run(status_line(post)))
    if post.excerpt:
        doc.add_paragraph().add_run(post.excerpt).italic = True
    _MarkdownWriter(doc).write(parse(post.content))


class _MarkdownWriter:
    """Walks markdown-it block tokens and emits Word paragraphs, lists, quotes and tables."""

    def __init__(self, doc: DocxDocument) -> None:
        self.doc = doc
        self.lists: list[str] = []  # "bullet" / "number", innermost last
        self.quote_depth = 0

    def write(self, tokens: list[Token]) -> None:
        index = 0
        while index < len(tokens):
            index = self._block(tokens, index)

    def _block(self, tokens: list[Token], index: int) -> int:
        token = tokens[index]
        kind = token.type
        if kind == "heading_open":
            # A post's own headings sit one level below its title.
            level = min(int(token.tag[1]) + 1, MAX_HEADING)
            self._inline(self.doc.add_heading(level=level), tokens[index + 1])
            return index + 3
        if kind == "paragraph_open":
            self._inline(self._paragraph(), tokens[index + 1])
            return index + 3
        if kind in {"bullet_list_open", "ordered_list_open"}:
            self.lists.append("bullet" if kind == "bullet_list_open" else "number")
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            self.lists.pop()
        elif kind == "blockquote_open":
            self.quote_depth += 1
        elif kind == "blockquote_close":
            self.quote_depth -= 1
        elif kind in {"fence", "code_block"}:
            self._code(token.content)
        elif kind == "hr":
            self._muted_rule()
        elif kind == "table_open":
            return self._table(tokens, index)
        return index + 1

    def _paragraph(self) -> Paragraph:
        if self.lists:
            base = "List Bullet" if self.lists[-1] == "bullet" else "List Number"
            depth = min(len(self.lists), 3)
            return self.doc.add_paragraph(style=base if depth == 1 else f"{base} {depth}")
        if self.quote_depth:
            return self.doc.add_paragraph(style="Quote")
        return self.doc.add_paragraph()

    def _inline(self, paragraph: Paragraph, token: Token) -> None:
        bold = italic = strike = False
        link: object | None = None  # the open <w:hyperlink>, if any
        for child in token.children or []:
            kind = child.type
            if kind in {"strong_open", "strong_close"}:
                bold = kind == "strong_open"
            elif kind in {"em_open", "em_close"}:
                italic = kind == "em_open"
            elif kind in {"s_open", "s_close"}:
                strike = kind == "s_open"
            elif kind == "link_open":
                href = safe_href(str(child.attrs.get("href", "")))
                link = _hyperlink(paragraph, href) if href else None
            elif kind == "link_close":
                link = None
            elif kind == "hardbreak":
                paragraph.add_run().add_break()
            elif kind == "softbreak":
                paragraph.add_run(" ")
            else:
                text = _inline_text(child)
                if not text:
                    continue
                run = paragraph.add_run(text)
                run.bold, run.italic = bold or None, italic or None
                run.font.strike = strike or None
                if kind == "code_inline":
                    run.font.name = MONO_FONT
                if kind == "image":
                    run.italic = True
                if link is not None:
                    run.font.color.rgb = LINK_COLOR
                    run.font.underline = True
                    link.append(run._r)  # type: ignore[attr-defined]  # move the run into the link

    def _code(self, content: str) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(12)
        _shade(paragraph, "F3F4F6")
        lines = content.rstrip("\n").split("\n")
        for number, line in enumerate(lines):
            run = paragraph.add_run(line)
            run.font.name = MONO_FONT
            run.font.size = Pt(9)
            if number < len(lines) - 1:
                run.add_break()

    def _muted_rule(self) -> None:
        paragraph = self.doc.add_paragraph()
        _muted(paragraph.add_run("* * *"))

    def _table(self, tokens: list[Token], index: int) -> int:
        rows: list[list[Token]] = []
        index += 1
        while tokens[index].type != "table_close":
            token = tokens[index]
            if token.type == "tr_open":
                rows.append([])
            elif token.type == "inline":
                rows[-1].append(token)
            index += 1
        columns = max((len(row) for row in rows), default=0)
        if columns:
            table = self.doc.add_table(rows=len(rows), cols=columns)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    paragraph = table.cell(r, c).paragraphs[0]
                    self._inline(paragraph, cell)
                    if r == 0:
                        for run in paragraph.runs:
                            run.bold = True
        return index + 1


def _inline_text(token: Token) -> str:
    if token.type == "image":
        alt = "".join(child.content for child in token.children or []) or "image"
        return f"[Image: {alt}]"
    return token.content


def _muted(run: object) -> None:
    run.font.color.rgb = MUTED_COLOR  # type: ignore[attr-defined]
    run.font.size = Pt(10)  # type: ignore[attr-defined]


def _hyperlink(paragraph: Paragraph, url: str) -> object:
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    paragraph._p.append(link)
    return link


def _internal_link(paragraph: Paragraph, anchor: str, text: str) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    paragraph._p.append(link)
    run = paragraph.add_run(text)
    run.font.color.rgb = LINK_COLOR
    link.append(run._r)


def _bookmark(paragraph: Paragraph, name: str, text: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), name.rsplit("-", 1)[-1])
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), name.rsplit("-", 1)[-1])
    paragraph._p.append(start)
    paragraph.add_run(text)
    paragraph._p.append(end)


def _shade(paragraph: Paragraph, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shading)
