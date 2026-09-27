"""Importing a file as a draft: formats, conversion, rejected files and size limits."""

import base64
import io
import zipfile
from collections.abc import AsyncIterator

from docx import Document
from httpx import AsyncClient

from tests.conftest import signed_in

IMPORT = "/api/v1/posts/import"
# A 1x1 transparent PNG, for a document with an image in it.
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


def docx_bytes(*, image: bool = False) -> bytes:
    doc = Document()
    doc.add_heading("Quarterly outlook", level=1)
    paragraph = doc.add_paragraph("Demand is ")
    paragraph.add_run("steady").bold = True
    doc.add_paragraph("Watch inventory", style="List Bullet")
    if image:
        doc.add_picture(io.BytesIO(PNG))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def zip_bytes(files: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return buffer.getvalue()


async def upload(
    client: AsyncClient, auth: dict[str, str], name: str, data: bytes
) -> tuple[int, dict[str, object]]:
    response = await client.post(IMPORT, files={"file": (name, data)}, headers=auth)
    body: dict[str, object] = response.json()
    return response.status_code, body


# ─── Accepted formats ─────────────────────────────────────


async def test_markdown_heading_becomes_the_title(client: AsyncClient) -> None:
    ada = await signed_in(client)

    status, body = await upload(client, ada, "notes.md", b"# Chip cycles\n\nSome *text*.\n")

    assert status == 200
    assert body == {"title": "Chip cycles", "content": "Some *text*.", "removed_images": 0}


async def test_text_file_is_titled_from_its_name_and_normalised(client: AsyncClient) -> None:
    ada = await signed_in(client)

    data = "\ufeffFirst line\r\nsecond line\r\n\r\n\r\n\r\nLast".encode()
    status, body = await upload(client, ada, "my_draft-post.txt", data)

    assert status == 200
    assert body["title"] == "My draft post"
    assert body["content"] == "First line\nsecond line\n\nLast"


async def test_docx_is_converted_to_markdown(client: AsyncClient) -> None:
    ada = await signed_in(client)

    status, body = await upload(client, ada, "outlook.docx", docx_bytes())

    assert status == 200
    assert body["title"] == "Quarterly outlook"
    assert "Demand is **steady**" in str(body["content"])
    assert "- Watch inventory" in str(body["content"])
    assert body["removed_images"] == 0


async def test_docx_images_are_dropped_and_counted(client: AsyncClient) -> None:
    ada = await signed_in(client)

    status, body = await upload(client, ada, "outlook.docx", docx_bytes(image=True))

    assert status == 200
    assert body["removed_images"] == 1
    assert "![" not in str(body["content"]) and "data:" not in str(body["content"])


async def test_import_saves_nothing(client: AsyncClient) -> None:
    ada = await signed_in(client)

    await upload(client, ada, "notes.md", b"# Title\n\nBody")

    response = await client.get("/api/v1/users/me/posts", headers=ada)
    assert response.json()["total"] == 0


# ─── Rejected files ───────────────────────────────────────


async def test_pdf_is_rejected_with_a_way_forward(client: AsyncClient) -> None:
    ada = await signed_in(client)

    status, body = await upload(client, ada, "report.pdf", b"%PDF-1.7 ...")

    assert status == 415
    assert ".docx" in str(body["detail"])


async def test_pdf_renamed_to_markdown_is_still_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)

    status, _ = await upload(client, ada, "report.md", b"%PDF-1.7 ...")

    assert status == 415


async def test_macro_documents_are_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)
    renamed = zip_bytes(
        {
            "[Content_Types].xml": b'<Types><Override ContentType="application/'
            b'vnd.ms-word.document.macroEnabled.main+xml"/></Types>',
            "word/document.xml": b"<w:document/>",
        }
    )

    docm_status, _ = await upload(client, ada, "plan.docm", docx_bytes())
    renamed_status, body = await upload(client, ada, "plan.docx", renamed)

    assert docm_status == 415
    assert renamed_status == 415 and "macros" in str(body["detail"])


async def test_binary_file_renamed_to_text_is_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)

    status, _ = await upload(client, ada, "image.md", PNG)

    assert status == 415


async def test_non_docx_zip_and_unknown_types_are_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)

    zip_status, _ = await upload(client, ada, "a.docx", zip_bytes({"hello.txt": b"hi"}))
    garbage_status, _ = await upload(client, ada, "a.docx", b"not a zip")
    html_status, _ = await upload(client, ada, "page.html", b"<p>Hi</p>")

    assert (zip_status, garbage_status, html_status) == (415, 415, 415)


async def test_archive_with_too_many_entries_is_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)
    files = {f"word/media/{i}.bin": b"" for i in range(1001)}
    files["word/document.xml"] = b"<w:document/>"

    status, body = await upload(client, ada, "big.docx", zip_bytes(files))

    assert status == 415 and "too large" in str(body["detail"])


async def test_archive_that_expands_too_far_is_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"<w:document/>")
        archive.writestr("word/padding.xml", b"\0" * (21 * 1024 * 1024))  # compresses to KBs

    status, _ = await upload(client, ada, "bomb.docx", buffer.getvalue())

    assert status == 415


async def test_empty_file_is_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)

    status, body = await upload(client, ada, "empty.md", b"  \n\n")

    assert status == 422 and "no text" in str(body["detail"])


async def test_file_over_one_megabyte_is_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)

    status, body = await upload(client, ada, "long.md", b"a" * (1024 * 1024 + 1))

    assert status == 413 and "1 MB" in str(body["detail"])


async def test_import_requires_sign_in(client: AsyncClient) -> None:
    response = await client.post(IMPORT, files={"file": ("a.md", b"hi")})

    assert response.status_code == 401


async def test_request_body_over_the_limit_is_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)
    body = b"a" * (2 * 1024 * 1024 + 1)

    response = await client.post(IMPORT, files={"file": ("a.md", body)}, headers=ada)

    assert response.status_code == 413
    assert response.headers["content-type"] == "application/problem+json"


async def test_streamed_body_over_the_limit_is_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)

    async def chunks() -> AsyncIterator[bytes]:  # no Content-Length: counted as it arrives
        yield b'--x\r\nContent-Disposition: form-data; name="file"; filename="a.md"\r\n\r\n'
        for _ in range(40):
            yield b"a" * (64 * 1024)

    response = await client.post(
        IMPORT,
        content=chunks(),
        headers={**ada, "Content-Type": "multipart/form-data; boundary=x"},
    )

    assert response.status_code == 413
