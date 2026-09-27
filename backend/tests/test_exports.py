"""Exporting posts as Word and PDF files: ownership, download headers and safe rendering."""

import io
import re
from datetime import UTC, datetime

from docx import Document
from httpx import AsyncClient

from app.rendering import pdf
from app.rendering.document import ExportDocument, ExportPost
from tests.conftest import create_post, signed_in

EXPORTS = "/api/v1/exports/posts"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def docx_text(data: bytes) -> str:
    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)


async def test_one_post_exports_as_docx(client: AsyncClient) -> None:
    ada = await signed_in(client)
    post = await create_post(
        client, ada, title="Reading a cloud bill", content="## Tag first\n\nThen **talk**."
    )

    response = await client.get(
        f"{EXPORTS}/{post['slug']}",
        params={"format": "docx"},
        headers={**ada, "Accept-Encoding": "gzip"},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_TYPE
    assert response.headers["content-disposition"] == (
        f'attachment; filename="nebula-{post["slug"]}.docx"'
    )
    assert response.headers["cache-control"] == "no-store"
    assert "content-encoding" not in response.headers  # already compressed, never gzipped
    text = docx_text(response.content)
    assert "Reading a cloud bill" in text and "Tag first" in text and "Then talk." in text
    assert "Draft, last edited" in text


async def test_all_posts_export_as_one_pdf(client: AsyncClient) -> None:
    ada = await signed_in(client)
    await create_post(client, ada, title="First", publish=True)
    await create_post(client, ada, title="Second")

    response = await client.get(EXPORTS, params={"format": "pdf"}, headers=ada)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert re.fullmatch(
        r'attachment; filename="nebula-posts-ada-\d{8}\.pdf"',
        response.headers["content-disposition"],
    )
    assert response.content.startswith(b"%PDF")


async def test_all_posts_docx_has_a_contents_list(client: AsyncClient) -> None:
    ada = await signed_in(client)
    await create_post(client, ada, title="First")
    await create_post(client, ada, title="Second")

    response = await client.get(EXPORTS, params={"format": "docx"}, headers=ada)

    text = docx_text(response.content)
    assert text.startswith("Posts by Ada")
    assert "Contents" in text
    assert text.index("First") < text.index("Second")


async def test_someone_elses_post_is_not_found(client: AsyncClient) -> None:
    ada = await signed_in(client, "ada")
    bob = await signed_in(client, "bob")
    post = await create_post(client, ada, publish=True)  # public, but still not Bob's to export

    response = await client.get(f"{EXPORTS}/{post['slug']}", params={"format": "pdf"}, headers=bob)

    assert response.status_code == 404


async def test_exporting_with_no_posts_is_not_found(client: AsyncClient) -> None:
    ada = await signed_in(client)

    response = await client.get(EXPORTS, params={"format": "docx"}, headers=ada)

    assert response.status_code == 404
    assert response.json()["detail"] == "You have no posts to export yet."


async def test_unknown_format_is_rejected(client: AsyncClient) -> None:
    ada = await signed_in(client)

    response = await client.get(EXPORTS, params={"format": "csv"}, headers=ada)

    assert response.status_code == 422


async def test_export_requires_sign_in(client: AsyncClient) -> None:
    response = await client.get(EXPORTS, params={"format": "pdf"})

    assert response.status_code == 401


async def test_exports_are_rate_limited(client: AsyncClient) -> None:
    ada = await signed_in(client)
    await create_post(client, ada)

    statuses = [
        (await client.get(EXPORTS, params={"format": "docx"}, headers=ada)).status_code
        for _ in range(11)
    ]

    assert statuses == [200] * 10 + [429]


def test_pdf_html_never_loads_images_or_unsafe_links() -> None:
    post = ExportPost(
        title="<script>alert(1)</script>",
        excerpt=None,
        content="![diagram](https://tracker.example/pixel.png)\n\n"
        "[click](javascript:alert(1)) and [docs](https://example.com)\n\n<b>raw</b>",
        tags=[],
        published=False,
        published_at=None,
        updated_at=datetime.now(UTC),
    )
    document = ExportDocument(
        title="t", author="a", generated_at=datetime.now(UTC), posts=[post], collection=False
    )

    html = pdf.to_html(document)

    assert "<img" not in html and "tracker.example" not in html
    assert 'href="javascript' not in html  # left as plain text, never a link
    assert 'href="https://example.com"' in html
    assert "<script>" not in html and "<b>raw</b>" not in html
